"""Сайт регистрации на «Долгожданный Трэш: Сумеречный лес»."""

import hashlib
import hmac
import io
import ipaddress
import logging
import logging.handlers
import os
import re
import secrets
import sqlite3
import threading
import time
import unicodedata
from collections import deque
from datetime import datetime, timedelta
from functools import wraps
from urllib.parse import quote

from flask import (
    Flask, Response, abort, flash, g, redirect, render_template, request,
    session, url_for,
)
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.environ.get("DB_PATH", os.path.join(BASE_DIR, "data", "registrations.db"))
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "")    # полный доступ
VIEWER_PASSWORD = os.environ.get("VIEWER_PASSWORD", "")  # только просмотр списка
SECRET_KEY = os.environ.get("SECRET_KEY", "")
PRODUCTION = os.environ.get("COOKIE_SECURE", "0") == "1"  # на сервере с HTTPS

if ADMIN_PASSWORD and ADMIN_PASSWORD == VIEWER_PASSWORD:
    raise SystemExit("ADMIN_PASSWORD и VIEWER_PASSWORD должны отличаться")
if PRODUCTION:
    problems = []
    if len(SECRET_KEY) < 32:
        problems.append("SECRET_KEY короче 32 символов")
    if len(ADMIN_PASSWORD) < 12:
        problems.append("ADMIN_PASSWORD короче 12 символов")
    if VIEWER_PASSWORD and len(VIEWER_PASSWORD) < 12:
        problems.append("VIEWER_PASSWORD короче 12 символов")
    if problems:
        raise SystemExit("Небезопасные настройки в .env: " + "; ".join(problems))

# задержки и лимиты для формы регистрации
SUBMIT_COOLDOWN = 5          # сек между отправками с одного устройства
AFTER_SUCCESS_COOLDOWN = 30  # сек после успешной заявки с этого устройства
MIN_FILL_SECONDS = 3         # быстрее человек форму не заполнит — это бот

EVENT = {
    "title": "Долгожданный треш!",
    "subtitle": "Сумеречный лес",
    "date": "1 ноября",
    "place": "Корпус УЛК",
    "schedule": [("15:00", "Начало"), ("18:00", "Квиз")],
}

COURSES = [
    "1 курс", "2 курс", "3 курс", "4 курс", "5 курс", "6 курс",
    "Магистратура, 1 курс", "Магистратура, 2 курс", "Аспирантура",
]
READINESS = [
    "Да, я готов(а) к сумеркам",
    "Родился(-ась) готовым(ой)",
    "Пока не знаю, но приду",
]

# принимаем и старые ссылки vk.com, но сохраняем всё как vk.ru
VK_RE = re.compile(r"^(?:https?://)?(?:www\.|m\.)?vk\.(?:ru|com)/([A-Za-z0-9_.]{2,64})/?$")

app = Flask(__name__)
app.config.update(
    SECRET_KEY=SECRET_KEY or secrets.token_hex(32),
    # __Host- : cookie только для этого домена и только по HTTPS
    SESSION_COOKIE_NAME="__Host-session" if PRODUCTION else "session",
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=PRODUCTION,
    PERMANENT_SESSION_LIFETIME=timedelta(hours=12),  # вход в админку живёт 12 часов
    MAX_CONTENT_LENGTH=32 * 1024,
    MAX_FORM_PARTS=50,
)
# ---------- логи ----------
# logs/access.log — каждый запрос одной строкой; logs/app.log — входы, лимиты, подозрительное.
# Файлы ротируются: по 5 МБ, хранится 5 штук. Если папка недоступна для записи — пишем только в консоль.
LOG_DIR = os.environ.get("LOG_DIR", os.path.join(BASE_DIR, "logs"))
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("lis_site")
access_log = logging.getLogger("lis_site.access")
access_log.propagate = False


def setup_file_logs():
    fmt = logging.Formatter("%(asctime)s %(levelname)s %(message)s")
    try:
        os.makedirs(LOG_DIR, exist_ok=True)
        for logger, name in ((log, "app.log"), (access_log, "access.log")):
            h = logging.handlers.RotatingFileHandler(
                os.path.join(LOG_DIR, name), maxBytes=5 * 1024 * 1024, backupCount=5, encoding="utf-8")
            h.setFormatter(logging.Formatter("%(asctime)s %(message)s") if logger is access_log else fmt)
            logger.addHandler(h)
    except OSError as e:
        access_log.addHandler(logging.StreamHandler())
        log.warning("не могу писать логи в %s (%s) — пишу только в консоль", LOG_DIR, e)


setup_file_logs()
if not SECRET_KEY:
    log.warning("SECRET_KEY не задан: сессии сбросятся при перезапуске")


# ---------- база данных ----------

def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH, timeout=10, isolation_level=None)
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    db = sqlite3.connect(DB_PATH)
    db.executescript(
        """
        PRAGMA journal_mode=WAL;
        CREATE TABLE IF NOT EXISTS registrations (
            id         INTEGER PRIMARY KEY AUTOINCREMENT,
            full_name  TEXT NOT NULL,
            faculty    TEXT NOT NULL,
            course     TEXT NOT NULL,
            vk_link    TEXT NOT NULL UNIQUE,
            source     TEXT NOT NULL,
            readiness  TEXT NOT NULL,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS settings (
            key   TEXT PRIMARY KEY,
            value TEXT NOT NULL
        );
        INSERT OR IGNORE INTO settings (key, value) VALUES ('registration_open', '1');
        """
    )
    # отметка «пришёл» — добавляем колонки и в уже существующую базу
    cols = {r[1] for r in db.execute("PRAGMA table_info(registrations)")}
    if "attended" not in cols:
        db.execute("ALTER TABLE registrations ADD COLUMN attended INTEGER NOT NULL DEFAULT 0")
    if "attended_at" not in cols:
        db.execute("ALTER TABLE registrations ADD COLUMN attended_at TEXT")
    db.commit()
    db.close()


def registration_open():
    row = get_db().execute("SELECT value FROM settings WHERE key='registration_open'").fetchone()
    return row is None or row["value"] == "1"


# ---------- защита ----------

class RateLimiter:
    """Скользящее окно в памяти процесса: не больше `limit` событий за `window` секунд на ключ.
    Потокобезопасно (gunicorn работает с потоками), устаревшие ключи периодически удаляются."""

    def __init__(self, limit, window):
        self.limit, self.window = limit, window
        self.hits, self.lock, self.swept = {}, threading.Lock(), time.monotonic()

    def hit(self, key):
        """Засчитать событие. Вернуть 0, если можно, иначе сколько секунд ждать."""
        now = time.monotonic()
        with self.lock:
            if now - self.swept > self.window:
                self.hits = {k: q for k, q in self.hits.items() if q and now - q[-1] < self.window}
                self.swept = now
            q = self.hits.setdefault(key, deque())
            while q and now - q[0] >= self.window:
                q.popleft()
            if len(q) >= self.limit:
                return int(self.window - (now - q[0])) + 1
            q.append(now)
            return 0


class Cooldown:
    """Пауза для ключа (устройства): до какого момента новые отправки запрещены."""

    def __init__(self):
        self.until, self.lock, self.swept = {}, threading.Lock(), time.monotonic()

    def left(self, key):
        with self.lock:
            return max(0, int(self.until.get(key, 0) - time.monotonic() + 0.999))

    def start(self, key, seconds):
        now = time.monotonic()
        with self.lock:
            if now - self.swept > 300:
                self.until = {k: t for k, t in self.until.items() if t > now}
                self.swept = now
            self.until[key] = max(self.until.get(key, 0), now + seconds)


# Лимиты по IP мягкие: в Wi-Fi университета сотни студентов выходят с одного адреса.
# От «прокликивания» с одного устройства защищает device_cooldown, а эти лимиты — от скриптов.
REGISTER_BURST = RateLimiter(20, 60)    # с одного IP: 20 отправок в минуту…
REGISTER_IP = RateLimiter(200, 600)     # …и 200 за 10 минут
LOGIN_IP = RateLimiter(10, 900)        # 10 попыток входа за 15 минут
device_cooldown = Cooldown()

# Откуда может прийти обратный прокси (nginx на том же сервере или Caddy в Docker-сети)
TRUSTED_PROXIES = [
    ipaddress.ip_network(n.strip())
    for n in os.environ.get("TRUSTED_PROXIES", "127.0.0.1/32,::1/128").split(",") if n.strip()
]


def from_trusted_proxy(addr):
    try:
        ip = ipaddress.ip_address(addr)
    except ValueError:
        return False
    return any(ip in net for net in TRUSTED_PROXIES)


def client_ip():
    # X-Real-IP выставляет прокси; верим ему только если запрос пришёл от самого прокси
    ip = request.remote_addr or "?"
    if from_trusted_proxy(ip):
        ip = request.headers.get("X-Real-IP", ip)
    return ip


def clean(value, limit=300):
    """Убрать управляющие и невидимые символы, схлопнуть пробелы."""
    value = unicodedata.normalize("NFC", value or "")[: limit * 2]
    value = "".join(ch for ch in value if unicodedata.category(ch) not in ("Cc", "Cf", "Co", "Cs"))
    return " ".join(value.split())


def csrf_token():
    if "csrf" not in session:
        session["csrf"] = secrets.token_hex(16)
    return session["csrf"]


app.jinja_env.globals["csrf_token"] = csrf_token


def csrf_ok():
    sent = request.form.get("csrf", "")
    return bool(sent) and hmac.compare_digest(sent, session.get("csrf", ""))


def check_csrf():
    if not csrf_ok():
        abort(400)


def role_fingerprint(role):
    """Отпечаток текущего пароля роли: сменили пароль в .env — старые входы перестают работать."""
    password = ADMIN_PASSWORD if role == "admin" else VIEWER_PASSWORD
    key = app.config["SECRET_KEY"].encode()
    return hmac.new(key, f"{role}:{password}".encode(), hashlib.sha256).hexdigest()[:24]


def current_role():
    role = session.get("role")
    if role in ("admin", "viewer") and hmac.compare_digest(session.get("pwv", ""), role_fingerprint(role)):
        return role
    return None


def login_required(view):
    """Любой вошедший: и админ, и наблюдатель."""
    @wraps(view)
    def wrapper(*args, **kwargs):
        if current_role() is None:
            session.pop("role", None)
            return redirect(url_for("admin_login"))
        return view(*args, **kwargs)
    return wrapper


def admin_required(view):
    """Только полный админ; наблюдателю — 403."""
    @wraps(view)
    def wrapper(*args, **kwargs):
        role = current_role()
        if role is None:
            session.pop("role", None)
            return redirect(url_for("admin_login"))
        if role != "admin":
            log.warning("наблюдатель пытался изменить данные: %s %s", request.path, client_ip())
            abort(403)
        return view(*args, **kwargs)
    return wrapper


def password_role(password):
    """Сверяем с обоими паролями всегда, чтобы время ответа не выдавало, какой из них подошёл."""
    is_admin = bool(ADMIN_PASSWORD) & hmac.compare_digest(password.encode(), ADMIN_PASSWORD.encode())
    is_viewer = bool(VIEWER_PASSWORD) & hmac.compare_digest(password.encode(), VIEWER_PASSWORD.encode())
    return "admin" if is_admin else "viewer" if is_viewer else None


CSP = "; ".join([
    "default-src 'self'",
    "script-src 'self'",
    "style-src 'self' https://fonts.googleapis.com",
    "font-src https://fonts.gstatic.com",
    "img-src 'self' data:",
    "connect-src 'self'",
    "form-action 'self'",
    "frame-ancestors 'none'",
    "base-uri 'none'",
    "object-src 'none'",
])


@app.before_request
def start_timer():
    g.t0 = time.perf_counter()


def log_safe(value, limit):
    # переводы строк и кавычки экранируем — в лог нельзя подсунуть поддельную строку
    return value[:limit].replace("\\", "\\\\").replace("\n", "\\n").replace("\r", "\\r").replace('"', '\\"')


@app.after_request
def log_request(resp):
    if request.path.startswith("/static/") or request.path == "/healthz":
        return resp
    ms = (time.perf_counter() - g.get("t0", time.perf_counter())) * 1000
    access_log.info('%s "%s %s" %s %s %.0fms "%s"',
                    client_ip(), request.method, log_safe(request.path, 300), resp.status_code,
                    resp.content_length if resp.content_length is not None else "-", ms,
                    log_safe(request.headers.get("User-Agent", "-"), 200))
    return resp


@app.after_request
def security_headers(resp):
    h = resp.headers
    h["Content-Security-Policy"] = CSP
    h["X-Content-Type-Options"] = "nosniff"
    h["X-Frame-Options"] = "DENY"
    h["Referrer-Policy"] = "strict-origin-when-cross-origin"
    h["Permissions-Policy"] = "camera=(), microphone=(), geolocation=(), payment=(), usb=()"
    h["Cross-Origin-Opener-Policy"] = "same-origin"
    if PRODUCTION:
        h["Strict-Transport-Security"] = "max-age=31536000"
    if not request.path.startswith("/static/"):
        # страницы с токенами и персональными данными не кешируем
        h["Cache-Control"] = "no-store"
    if request.path.startswith("/admin"):
        h["Referrer-Policy"] = "no-referrer"
        h["X-Robots-Tag"] = "noindex, nofollow"
    return resp


ERRORS = {  # код: (подзаголовок, текст)
    400: ("Что-то пошло не так", "Запрос не понят. Обновите страницу и попробуйте ещё раз."),
    403: ("Сюда нельзя", "У этого входа нет прав на такое действие."),
    404: ("Ты заблудился в лесу", "Такой тропы здесь нет. Возвращайся к костру — на главную."),
    405: ("Не той тропой", "Так к этой странице обращаться нельзя."),
    413: ("Слишком тяжело", "Запрос слишком большой."),
    429: ("Притормози", "Слишком много запросов. Подожди немного и попробуй снова."),
    500: ("Лес шумит", "Что-то сломалось на нашей стороне. Попробуй через пару минут."),
}


@app.errorhandler(400)
@app.errorhandler(403)
@app.errorhandler(404)
@app.errorhandler(405)
@app.errorhandler(413)
@app.errorhandler(429)
@app.errorhandler(500)
def error_page(err):
    code = getattr(err, "code", 500) or 500
    kicker, message = ERRORS.get(code, ERRORS[500])
    return render_template("error.html", code=code, kicker=kicker, message=message), code


@app.get("/healthz")
def healthz():
    # для проверки здоровья контейнера; в логи не пишется
    get_db().execute("SELECT 1")
    return "ok"


@app.get("/favicon.ico")
def favicon():
    # браузеры просят значок по этому адресу сами, даже без <link rel="icon">
    return app.send_static_file("favicon.ico")


@app.get("/robots.txt")
def robots():
    return Response("User-agent: *\nDisallow: /admin\n", mimetype="text/plain")


# ---------- публичная часть ----------

def render_index(errors=None, form=None, status=200, retry_after=0):
    resp = app.make_response((render_template(
        "index.html",
        event=EVENT,
        courses=COURSES,
        readiness=READINESS,
        is_open=registration_open(),
        errors=errors or {},
        form=form or {},
        retry_after=retry_after,
    ), status))
    if retry_after:
        resp.headers["Retry-After"] = str(retry_after)
    return resp


def start_form():
    """Метки устройства и времени открытия формы — для паузы между отправками и ловушки на ботов."""
    session.setdefault("did", secrets.token_hex(12))
    session["form_at"] = time.time()


@app.get("/")
def index():
    start_form()
    return render_index(retry_after=device_cooldown.left(session["did"]))


@app.post("/register")
def register():
    limits = {"full_name": 120, "faculty": 120, "course": 60, "vk_link": 100, "source": 200, "readiness": 60}
    form = {k: clean(request.form.get(k), n) for k, n in limits.items()}
    ip = client_ip()

    # 1. лимиты по IP — режем поток раньше всего остального
    wait = REGISTER_BURST.hit(ip) or REGISTER_IP.hit(ip)
    if wait:
        log.warning("лимит регистраций по IP %s", ip)
        return render_index({"_form": f"Слишком много заявок с этой сети. Попробуйте через {wait} с."}, form, 429, wait)

    # 2. CSRF + страница должна быть открыта на этом устройстве (есть метка устройства)
    did = session.get("did")
    if not csrf_ok() or not did:
        start_form()
        return render_index({"_form": "Страница устарела — проверьте данные и отправьте ещё раз."}, form, 400)

    # 3. пауза между отправками с одного устройства
    wait = device_cooldown.left(did)
    if wait:
        return render_index({"_form": f"Подождите {wait} с перед повторной отправкой."}, form, 429, wait)
    device_cooldown.start(did, SUBMIT_COOLDOWN)

    # 4. боты: заполненное скрытое поле или форма отправлена мгновенно
    if request.form.get("website"):
        log.info("honeypot сработал: %s", ip)
        return redirect(url_for("success"))
    if time.time() - session.get("form_at", 0) < MIN_FILL_SECONDS:
        session["form_at"] = time.time()
        return render_index({"_form": "Слишком быстро. Проверьте данные и отправьте ещё раз."}, form, 400, SUBMIT_COOLDOWN)

    errors = {}
    if len(form["full_name"].split()) < 2 or len(form["full_name"]) > 120:
        errors["full_name"] = "Укажите фамилию и имя (можно с отчеством)."
    if not 2 <= len(form["faculty"]) <= 120:
        errors["faculty"] = "Укажите факультет."
    if form["course"] not in COURSES:
        errors["course"] = "Выберите курс."
    m = VK_RE.match(form["vk_link"])
    if not m:
        errors["vk_link"] = "Нужна ссылка вида vk.ru/username."
    if not 2 <= len(form["source"]) <= 200:
        errors["source"] = "Напишите, откуда узнали о мероприятии."
    if form["readiness"] not in READINESS:
        errors["readiness"] = "Выберите вариант."
    if errors:
        return render_index(errors, form, 400, SUBMIT_COOLDOWN)

    vk_link = f"https://vk.ru/{m.group(1).lower()}"

    db = get_db()
    if not registration_open():
        return render_index({"_form": "Увы, регистрация уже закрыта."}, form, 409)
    try:
        db.execute(
            "INSERT INTO registrations (full_name, faculty, course, vk_link, source, readiness, created_at)"
            " VALUES (?, ?, ?, ?, ?, ?, ?)",
            (form["full_name"], form["faculty"], form["course"], vk_link, form["source"],
             form["readiness"], datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        )
    except sqlite3.IntegrityError:
        return render_index({"vk_link": "С этой страницей ВК уже зарегистрировались."}, form, 409, SUBMIT_COOLDOWN)

    device_cooldown.start(did, AFTER_SUCCESS_COOLDOWN)
    session["form_at"] = time.time()
    return redirect(url_for("success"))


@app.get("/success")
def success():
    return render_template("success.html", event=EVENT)


# ---------- админка ----------

@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    error = None
    if request.method == "POST":
        check_csrf()
        ip = client_ip()
        if wait := LOGIN_IP.hit(ip):
            log.warning("лимит входов в админку: %s", ip)
            error = f"Слишком много попыток. Подождите {wait // 60 + 1} мин."
            return render_template("login.html", error=error), 429, {"Retry-After": str(wait)}
        if role := password_role(request.form.get("password", "")[:200]):
            session.clear()  # новая сессия при входе — защита от фиксации сессии
            session.permanent = True
            session["role"], session["pwv"] = role, role_fingerprint(role)
            log.info("вход в админку: %s (%s)", role, ip)
            return redirect(url_for("admin"))
        log.warning("неверный пароль админки: %s", ip)
        error = "Неверный пароль."
    return render_template("login.html", error=error)


@app.post("/admin/logout")
def admin_logout():
    check_csrf()
    session.clear()
    return redirect(url_for("index"))


@app.get("/admin")
@login_required
def admin():
    rows = get_db().execute("SELECT * FROM registrations ORDER BY id").fetchall()
    came = sum(1 for r in rows if r["attended"])
    return render_template("admin.html", rows=rows, came=came, is_open=registration_open(), event=EVENT,
                           can_edit=current_role() == "admin")


@app.post("/admin/toggle")
@admin_required
def admin_toggle():
    check_csrf()
    new = "0" if registration_open() else "1"
    get_db().execute("UPDATE settings SET value=? WHERE key='registration_open'", (new,))
    flash("Регистрация открыта." if new == "1" else "Регистрация закрыта.")
    return redirect(url_for("admin"))


@app.post("/admin/delete/<int:reg_id>")
@admin_required
def admin_delete(reg_id):
    check_csrf()
    get_db().execute("DELETE FROM registrations WHERE id=?", (reg_id,))
    flash("Запись удалена.")
    return redirect(url_for("admin"))


@app.post("/admin/attend/<int:reg_id>")
@admin_required
def admin_attend(reg_id):
    check_csrf()
    attended = request.form.get("attended") == "1"
    db = get_db()
    db.execute(
        "UPDATE registrations SET attended=?, attended_at=? WHERE id=?",
        (int(attended), datetime.now().strftime("%Y-%m-%d %H:%M:%S") if attended else None, reg_id),
    )
    came = db.execute("SELECT COUNT(*) FROM registrations WHERE attended=1").fetchone()[0]
    if request.headers.get("X-Requested-With") == "fetch":
        return {"ok": True, "attended": attended, "came": came}
    return redirect(url_for("admin"))


@app.get("/admin/export.xlsx")
@admin_required
def admin_export():
    rows = get_db().execute("SELECT * FROM registrations ORDER BY id").fetchall()
    data = build_xlsx(rows)
    name = f"участники_{datetime.now():%Y-%m-%d}.xlsx"
    return Response(
        data,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=participants.xlsx; filename*=UTF-8''{quote(name)}"},
    )


def as_text(cell):
    """Защита от Excel-инъекции: то, что ввёл участник, всегда текст, даже если начинается с «=»."""
    if isinstance(cell.value, str):
        cell.data_type = "s"
    return cell


def build_xlsx(rows):
    wb = Workbook()
    orange = PatternFill("solid", fgColor="E94B26")
    head_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(vertical="top", wrap_text=True)

    # --- лист со всеми участниками ---
    ws = wb.active
    ws.title = "Участники"
    headers = ["№", "ФИО", "Факультет", "Курс", "Страница ВК", "Откуда узнал(а)",
               "Готовность", "Пришёл", "Время отметки", "Дата регистрации"]
    widths = [6, 34, 24, 22, 32, 34, 30, 10, 18, 18]
    ws.append(headers)
    for col, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=col)
        c.fill, c.font, c.border = orange, head_font, border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 30

    as_dt = lambda v: datetime.strptime(v, "%Y-%m-%d %H:%M:%S") if v else None
    for i, r in enumerate(rows, 1):
        ws.append([i, r["full_name"], r["faculty"], r["course"], r["vk_link"], r["source"],
                   r["readiness"], "Да" if r["attended"] else "Нет",
                   as_dt(r["attended_at"]), as_dt(r["created_at"])])
        row = i + 1
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=row, column=col)
            c.border, c.alignment = border, wrap
            if 2 <= col <= 7:
                as_text(c)
        link = ws.cell(row=row, column=5)
        link.hyperlink, link.font = r["vk_link"], Font(color="0563C1", underline="single")
        for col in (9, 10):
            ws.cell(row=row, column=col).number_format = "DD.MM.YYYY HH:MM"
        ws.cell(row=row, column=8).alignment = Alignment(horizontal="center", vertical="top")

    last = len(rows) + 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    if rows:
        green, grey = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="F2F2F2")
        ws.conditional_formatting.add(f"H2:H{last}", CellIsRule(operator="equal", formula=['"Да"'], fill=green, font=Font(color="006100", bold=True)))
        ws.conditional_formatting.add(f"H2:H{last}", CellIsRule(operator="equal", formula=['"Нет"'], fill=grey))

    # --- сводка (формулы пересчитаются, если поменять «Пришёл» в таблице) ---
    st = wb.create_sheet("Сводка")
    st.column_dimensions["A"].width = 36
    for col in "BCD":
        st.column_dimensions[col].width = 16
    rng = lambda col: f"'Участники'!${col}$2:${col}${max(last, 2)}"
    st["A1"] = f"{EVENT['title']} {EVENT['subtitle']} — {EVENT['date']}, {EVENT['place']}"
    st["A1"].font = Font(bold=True, size=14)
    totals = [
        ("Зарегистрировались", f"=COUNTA({rng('B')})"),
        ("Пришли", f'=COUNTIF({rng("H")},"Да")'),
        ("Не пришли", "=B3-B4"),
        ("Явка", "=IF(B3=0,0,B4/B3)"),
    ]
    for n, (label, formula) in enumerate(totals, 3):
        st.cell(row=n, column=1, value=label).font = Font(bold=True)
        st.cell(row=n, column=2, value=formula)
    st["B6"].number_format = "0%"

    row = 8
    for title, col, key in (("По факультетам", "C", "faculty"), ("По курсам", "D", "course"),
                            ("Готовность к вечеру", "G", "readiness")):
        for n, h in enumerate((title, "Записались", "Пришли", "Явка"), 1):
            c = st.cell(row=row, column=n, value=h)
            c.fill, c.font, c.border = orange, head_font, border
        values = sorted({r[key] for r in rows}, key=lambda v: (COURSES.index(v) if v in COURSES else 99, v))
        for v in values:
            row += 1
            as_text(st.cell(row=row, column=1, value=v))
            st.cell(row=row, column=2, value=f"=COUNTIF({rng(col)},A{row})")
            st.cell(row=row, column=3, value=f'=COUNTIFS({rng(col)},A{row},{rng("H")},"Да")')
            st.cell(row=row, column=4, value=f"=IF(B{row}=0,0,C{row}/B{row})").number_format = "0%"
            for n in range(1, 5):
                st.cell(row=row, column=n).border = border
        row += 3

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


init_db()

if __name__ == "__main__":
    # локальный запуск; отладчик включается только явно: FLASK_DEBUG=1
    app.run(host="127.0.0.1", port=5050, debug=os.environ.get("FLASK_DEBUG") == "1")
